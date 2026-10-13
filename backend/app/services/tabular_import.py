"""Canonical JSON and bounded, non-executing XLSX worksheet ingestion."""
from __future__ import annotations

import io
import json
import zipfile

import pandas as pd


def json_frame(raw):
    def reject_constant(value):
        raise ValueError('JSON numbers must be finite.')
    def unique_object(pairs):
        record={}
        for key,value in pairs:
            if key in record:raise ValueError('Duplicate JSON fields are not permitted.')
            record[key]=value
        return record
    value=json.loads(raw,parse_constant=reject_constant,object_pairs_hook=unique_object)
    if not isinstance(value,list) or not value or any(not isinstance(row,dict) for row in value):
        raise ValueError('JSON must be a nonempty array of flat study records.')
    return pd.DataFrame(value)


def xlsx_frame(raw,config):
    from defusedxml.ElementTree import fromstring
    from openpyxl import load_workbook

    expanded_limit=min(100*1024*1024,config.max_upload_bytes*10)
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        entries=archive.infolist()
        if any(entry.filename.lower().endswith('vbaproject.bin') for entry in entries):
            raise ValueError('Macro-enabled workbooks are not supported.')
        if len(entries)>1000 or sum(entry.file_size for entry in entries)>expanded_limit:
            raise ValueError('Expanded workbook exceeds the import limit.')
        if any(entry.flag_bits & 1 for entry in entries):
            raise ValueError('Encrypted workbooks are not supported.')
        for entry in entries:
            if entry.filename.endswith(('.xml','.rels')):
                # Reject DTD/entities before the workbook library processes XML.
                fromstring(archive.read(entry),forbid_dtd=True,forbid_entities=True,forbid_external=True)
    workbook=load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False)
    try:
        if not workbook.worksheets:raise ValueError('Workbook has no worksheet.')
        sheet=workbook.worksheets[0]
        if sheet.max_column is not None and sheet.max_column>100:raise ValueError('At most 100 columns are supported.')
        if sheet.max_row is not None and sheet.max_row>config.max_dataset_rows+1:raise ValueError('Workbook row count exceeds the import limit.')
        iterator=sheet.iter_rows(values_only=True)
        header=next(iterator,())
        header=[str(value).strip() if value is not None else '' for value in header]
        if not header or len(header)>100 or any(not value for value in header) or len(set(header))!=len(header):
            raise ValueError('Column names must be unique and non-empty.')
        rows=[]
        for values in iterator:
            if not any(value is not None for value in values):continue
            if len(values)!=len(header):raise ValueError('Every worksheet row must match the header width.')
            rows.append(values)
            if len(rows)>config.max_dataset_rows or len(rows)*len(header)>2_000_000:
                raise ValueError('Worksheet exceeds bounded row/cell limits.')
        return pd.DataFrame(rows,columns=header)
    finally:
        workbook.close()
